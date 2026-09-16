from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROLE_ZH = {
    "Mechanism Reviewer": "创新性与机制审稿人",
    "Clinical Validity Reviewer": "临床与外部有效性审稿人",
    "Statistical Reviewer": "统计严谨性审稿人",
    "Reproducibility Reviewer": "复现性与伦理审稿人",
    "Figures & Tables Reviewer": "图表与写作审稿人",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_labelled_paragraph(doc: Document, heading: str, body: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(f"{heading}。").bold = True
    p.add_run(body)


def markdown(payload: dict) -> str:
    lines = ["# DCS-OPCT v11 BSPC 最终多视角审查报告", ""]
    level = payload["editor_letter"]["revision_level"]
    lines += [f"## 编辑信 — {level}", ""]
    for paragraph in payload["editor_letter"]["paragraphs"]:
        lines += [paragraph, ""]
    for idx, reviewer in enumerate(payload["reviewers"], 1):
        role = ROLE_ZH.get(reviewer["role_label"], reviewer["role_label"])
        lines += [f"## 审稿人 {idx} — {role}", "", "### 整体点评", "", reviewer["overall_paragraph"], ""]
        lines += ["### 主要问题", ""]
        if reviewer["major_concerns"]:
            for number, concern in enumerate(reviewer["major_concerns"], 1):
                lines += [
                    f"{number}. **{concern['heading']}。** {concern['body']} 证据位置：{concern['evidence_pointer']}。",
                    "",
                ]
        else:
            lines += ["无新增主要问题。", ""]
        lines += ["### 次要问题", ""]
        if reviewer["minor_concerns"]:
            for number, concern in enumerate(reviewer["minor_concerns"], 1):
                lines += [f"{number}. **{concern['heading']}。** {concern['body']}", ""]
        else:
            lines += ["无。", ""]
        lines += [f"**置信度：** {reviewer['overall_confidence']}", ""]

    lines += ["## 跨审稿意见汇总", ""]
    for item in payload["consensus_block"]:
        raised = "、".join(item["raised_by"])
        lines += [
            f"**{item['id']} — {item['description']}** | 提出者：{raised} | 轴：{item['axis']} | 严重程度：{item['severity']}",
            "",
            item["rationale"],
            "",
        ]

    table = payload["revision_task_table"]
    lines += ["## 修订任务表", "", "| " + " | ".join(table["headers"]) + " |"]
    lines += ["|" + "|".join(["---"] * len(table["headers"])) + "|"]
    for row in table["rows"]:
        lines.append("| " + " | ".join(str(value).replace("|", "\\|") for value in row) + " |")
    lines += ["", "由 nature-review-studio 生成 · 2026-09-14 · dcs-opct-v11-bspc-final-rereview", ""]
    return "\n".join(lines)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    for style_name, size, bold in (
        ("Normal", 10.25, False),
        ("Title", 16, True),
        ("Heading 1", 13, True),
        ("Heading 2", 11, True),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_p_pr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15


def build_docx(payload: dict, out: Path) -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DCS OPCT v11 BSPC 最终多视角审查报告")
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    level = payload["editor_letter"]["revision_level"]
    doc.add_paragraph(f"编辑信  {level}", style="Heading 1")
    for paragraph in payload["editor_letter"]["paragraphs"]:
        doc.add_paragraph(paragraph)

    for idx, reviewer in enumerate(payload["reviewers"], 1):
        role = ROLE_ZH.get(reviewer["role_label"], reviewer["role_label"])
        doc.add_paragraph(f"审稿人 {idx}  {role}", style="Heading 1")
        doc.add_paragraph("整体点评", style="Heading 2")
        doc.add_paragraph(reviewer["overall_paragraph"])
        doc.add_paragraph("主要问题", style="Heading 2")
        if reviewer["major_concerns"]:
            for concern in reviewer["major_concerns"]:
                body = f"{concern['body']} 证据位置：{concern['evidence_pointer']}。"
                add_labelled_paragraph(doc, concern["heading"], body)
        else:
            doc.add_paragraph("无新增主要问题。")
        doc.add_paragraph("次要问题", style="Heading 2")
        if reviewer["minor_concerns"]:
            for concern in reviewer["minor_concerns"]:
                add_labelled_paragraph(doc, concern["heading"], concern["body"])
        else:
            doc.add_paragraph("无。")
        p = doc.add_paragraph()
        p.add_run("置信度：").bold = True
        p.add_run(reviewer["overall_confidence"])

    doc.add_paragraph("跨审稿意见汇总", style="Heading 1")
    for item in payload["consensus_block"]:
        raised = "、".join(item["raised_by"])
        p = doc.add_paragraph()
        p.add_run(f"{item['id']}  {item['description']}").bold = True
        p.add_run(f" | 提出者：{raised} | 轴：{item['axis']} | 严重程度：{item['severity']}")
        doc.add_paragraph(item["rationale"])

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Cm(29.7)
    landscape.page_height = Cm(21)
    landscape.top_margin = Cm(1.8)
    landscape.bottom_margin = Cm(1.8)
    landscape.left_margin = Cm(1.5)
    landscape.right_margin = Cm(1.5)
    doc.add_paragraph("修订任务表", style="Heading 1")
    table_data = payload["revision_task_table"]
    table = doc.add_table(rows=1, cols=len(table_data["headers"]))
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_header(table.rows[0])
    prevent_split(table.rows[0])
    for idx, header in enumerate(table_data["headers"]):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    for row_index, values in enumerate(table_data["rows"], 1):
        row = table.add_row()
        prevent_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
                    run.font.name = "Calibri"
                    r_pr = run._element.get_or_add_rPr()
                    r_pr.get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")

    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "由 nature-review-studio 生成 · 2026-09-14 · dcs-opct-v11-bspc-final-rereview"
        for run in footer.runs:
            run.font.size = Pt(9)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            r_pr = run._element.get_or_add_rPr()
            r_pr.get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--payload", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    payload = json.loads(args.payload.read_text(encoding="utf-8"))
    build_docx(payload, args.out)
    md_out = args.out.with_suffix(".md")
    md_out.write_text(markdown(payload), encoding="utf-8")
    print(args.out.resolve())
    print(md_out.resolve())


if __name__ == "__main__":
    main()
