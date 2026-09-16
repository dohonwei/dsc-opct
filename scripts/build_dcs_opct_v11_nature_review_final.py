from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/dcs_opct_v11_nature_review_20260914_final"
STEM = "review_dcs_opct_v11_20260914_final"
DOCX = OUT / f"{STEM}.docx"
MARKDOWN = OUT / f"{STEM}.md"

EDITOR = [
    (
        "Major revision。当前版本已经从对比实验稿推进为机制、风险与选择性校准相连的完整方法论文，"
        "并以原始数据先分区实验关闭了最直接的训练数据重叠质疑。方法方程、端点不确定性、"
        "信息访问边界、失败回退和复现证据均已达到较高完整度。"
    ),
    (
        "现阶段唯一仍具有实质性的科学缺口，是尚无一个预先保留且通过全部冻结门控后产生非 identity action 的外部数据集。"
        "因此稿件可以支持开发域内的选择性配置级审计风险校准及完整分区下的 fail-closed 执行，"
        "但不能支持前瞻外部有效性或普遍安全迁移。若继续以当前收紧后的结论投稿 BSPC，风险已被诚实界定；"
        "若要把中心结论升级为有效的安全跨域迁移，则仍需完成冻结 one-shot 外部确认。"
    ),
]

REVIEWERS = [
    {
        "role": "审稿人 1 方法创新与机器学习",
        "overall": (
            "稿件的创新不再依赖 CORAL、quantile mapping 或 logistic calibration 单个组件，"
            "而在于把反事实身份利用端点、配置级风险预测、适用性筛查、有限标签证书和返回原概率动作连接为可审计流程。"
            "最新独立性阶梯使贡献边界清楚，但也显示正向 release 随独立性增强而收缩。"
        ),
        "major": [
            (
                "中心贡献与最强证据层级",
                "完全参与者和训练行分离的 EPPVR 实验在冻结适用性门控处停止，未形成正向校准证据。稿件已在 Abstract、Discussion、Limitations 和 Conclusion 中明确这一点，因此当前可辩护的贡献是选择性开发域校准与 fail-closed 执行，而不是已验证的通用安全迁移。该边界必须在投稿信和所有宣传性文本中保持不变。",
                "Abstract；Results 的 independence ladder；Discussion；Conclusion",
            ),
            (
                "方法原语与系统能力的区分",
                "能力矩阵已明确 DCS-OPCT 不是新的 alignment、calibration 或 certification 原语，并以端到端能力组合定位创新。OPCT 软目标、正斜率约束、正则项和参数范围均已公式化；两个变换也已改称 separately constructed，并明确其一致性不是统计独立确认。该问题已充分解决。",
                "Table 1；Methods Stage III，Eqs. 6--7",
            ),
        ],
        "minor": [
            "投稿题目、Highlights 和 cover letter 应继续使用 audit-risk calibration 或 risk-controlled release，不再使用 safe cross-domain transfer。",
            "避免把 return-original 写成 successful adaptation；现有图注已经采用正确语义。",
        ],
        "confidence": "高",
    },
    {
        "role": "审稿人 2 统计与实验设计",
        "overall": (
            "作者已增加完整剂量 cluster、较高层聚合敏感性、端点 bootstrap、匹配预算比较和训练行完全分离实验，"
            "统计透明度明显改善。剩余限制主要来自独立数据集数量和单次 15/15 原始分区，而不是检验实现错误。"
        ),
        "major": [
            (
                "独立域样本量仍为两个",
                "Stage II 冻结风险模型的真正独立训练域只有 DEAP 和 MAHNOB-HCI。480 个配置不能替代 dataset-level n=2。稿件已在 Limitations 中明确这一层级，并报告七域后验重拟合未通过预设门控；因此现有结果应被解释为机制驱动的开发证据，而非稳定的跨域参数估计。增加新的预先保留兼容域仍是根本解决办法。",
                "Methods Stage II；Results risk transfer；Limitations 第一项",
            ),
            (
                "端点测量误差与阈值意义",
                "交叉 participant/stimulus bootstrap 显示几乎所有源配置区间跨越 0.02，且无 LCB event。作者已把 0.02 限定为冻结的工程阈值而非生理状态，并用 soft-label 分析检验动作方向。该处理足以回答不确定性意见，但不支持阈值最优性或机制系数稳定性。",
                "Methods endpoint sensitivity；Supplementary Tables S24--S25；Limitations 第三项",
            ),
            (
                "完整分区结果只有一次人口划分",
                "15/15 parity split 严格关闭了跨人口训练重叠，21/21 验证也确认实现正确；但它仍是单数据集单划分。当前文本已将其限定为 sensitivity，而不是独立效应估计。除非重新进行多次预先规定的人口划分，否则不宜给出该门控停止概率或普遍稳定性结论。",
                "Methods raw-partition-first sensitivity；Figure 3；Supplementary Table S23",
            ),
        ],
        "minor": [
            "100 个 audit orderings 应始终称为固定数据条件下的 allocation sensitivity，不能称为独立重复。",
            "主文已移除 post-hoc 20% 单元的选择细节，完整检验应继续保留在 Supplementary。",
        ],
        "confidence": "高",
    },
    {
        "role": "审稿人 3 生物医学信号与外部有效性",
        "overall": (
            "EPPVR 的设备、ADS1299 前端、右耳参考与 ground、基线、采样结构、伦理审批和书面同意已得到报告，"
            "并且论文明确研究对象是 benchmark configuration audit risk，而非直接改善患者或用户的情绪预测。"
            "真正未闭合的是成功的前瞻外部有效释放。"
        ),
        "major": [
            (
                "缺少成功的预先冻结外部 release",
                "AVDOS-VR 和 FACED 是 post-access exploratory，EEGEmotions-27 在无标签适用性门控处停止，AMIGOS 与 Emognition 又因数据访问尚未执行。故当前没有一个未参与开发的兼容域在冻结规则下通过全部门控并产生有效非 identity action。该缺口不能通过补写文本消除，只能通过取得预先保留数据并一次性执行已锁定协议解决。",
                "Methods evidence roles；Results external boundaries；Limitations 最后一项",
            ),
            (
                "跨数据集标签构念差异",
                "正文已说明不同域使用 participant-rated dimensions 或 category contrasts，Supplementary Table S1 给出 endpoint ontology 和排除规则。DCS-OPCT 迁移的是局部定义后形成的配置级 material-optimism event，而不是直接合并情绪标签。该限制已经透明化，但仍要求未来外部确认使用兼容且预先声明的 endpoint。",
                "Methods Study endpoint and evidence roles；Supplementary Table S1",
            ),
        ],
        "minor": [
            "正式投稿前应依据伦理批件确认伦理委员会官方英文名称。",
            "channels 5--9 的确切顺序只有在采集记录可核实时再补充，不能根据通道数量推测。",
        ],
        "confidence": "高",
    },
    {
        "role": "审稿人 4 复现性与图表呈现",
        "overall": (
            "稿件具有异常完整的工程审计链：GPU 运行、进度条、固定随机种子、冻结哈希、失败记录、"
            "9,760 行拟合溯源、独立验证和 252 文件复现清单。新独立性阶梯图清楚展示证据收缩，"
            "LaTeX 日志和关键页面也已通过版面检查。"
        ),
        "major": [
            (
                "匿名代码入口尚未填入",
                "内部复现链已经通过 23/23 inventory checks 和 14/14 reviewer-closure checks，但稿件仍只有 repository URL 或 DOI 的作者 TODO。审稿人无法从本地路径访问这些证据。正式投稿前必须建立匿名只读仓库或归档 DOI，并确认受限数据只发布许可允许的派生物。",
                "Data and code availability；README；Supplementary Reproducibility",
            ),
            (
                "主文信息密度仍高",
                "主文曾包含大量开发史、比较器和边界分析。最终修订已将重复的 cross-fit stress-test 和 same-ceiling comparator 图移入 42 页 Supplementary，主文由 43 页降至 42 页，并保留三阶段主线、独立性阶梯、预算风险和外部边界。该项已在不改变结果的前提下完成；进一步压缩属于编辑性优化。",
                "Results 全节；Figures 1--10",
            ),
        ],
        "minor": [
            "补充材料 Table S23 字体较小但未裁切；最终双栏排版时需重新检查可读性。",
            "Funding statement 仍需作者确认；若无专项资助，应给出明确 no specific funding 声明。",
        ],
        "confidence": "高",
    },
]

CONSENSUS = [
    (
        "Cc.1",
        "外部有效性仍是唯一实质性科学缺口",
        "审稿人 1、2、3",
        "clinical-validity",
        "完整数据隔离验证了 fail-closed 执行，却没有产生正向 release。要升级为前瞻有效跨域框架，仍需预先保留的兼容外部域。",
    ),
    (
        "Cc.2",
        "结论必须服从独立性阶梯",
        "审稿人 1、2、3",
        "claim-moderation",
        "正向证据从 transductive 到 raw-test-group separation 逐步收缩，最强分离层级为 non-intervention。当前稿件已经正确限定，投稿材料不得放宽。",
    ),
    (
        "Cc.3",
        "投稿前仍需公开入口和作者确认",
        "审稿人 3、4",
        "reproducibility",
        "匿名仓库或 DOI、Funding statement、伦理英文名称和可核实通道顺序属于作者依赖项目，不能由计算分析替代。",
    ),
]

TASKS = [
    ("T1", "R1/R2/R3", "完成预先冻结的兼容外部 one-shot", "按既有 lock 一次执行并保留全部结果", "TODO_EXPERIMENT", "AMIGOS 或 Emognition 授权数据", "非 identity effectiveness 或诚实 failure receipt", "Yes"),
    ("T2", "R1/R2/R3", "保持收紧后的 claim boundary", "同步 title abstract highlights cover letter", "DONE", "无需新数据", "不宣称普遍安全或外部有效性", "No"),
    ("T3", "R1", "补全 OPCT 方程和 witness 关系", "公式化并声明非统计独立", "DONE", "无需新数据", "Methods Eqs. 6--7", "No"),
    ("T4", "R2", "关闭训练数据重叠质疑", "15/15 raw-partition-first GPU 重建", "DONE", "EPPVR 原始数据", "9,600 fits 和 21/21 validation", "No"),
    ("T5", "R2", "传播 0.02 endpoint 测量误差", "paired crossed bootstrap 与 soft labels", "DONE", "现有预测", "不确定性表与限定结论", "No"),
    ("T6", "R4", "发布匿名复现入口", "只读仓库或归档 DOI", "TODO_AUTHOR_CONFIRM", "公开位置与许可确认", "Data and code availability URL", "Yes"),
    ("T7", "R3/R4", "确认投稿元数据", "核对伦理英文名与 Funding", "TODO_AUTHOR_CONFIRM", "伦理批件和资助信息", "最终 ethics 和 funding statements", "Yes"),
    ("T8", "R4", "压缩主文信息负载", "将重复诊断图移至补充材料，保留核心图", "DONE", "无需新数据", "更紧凑的投稿稿", "No"),
]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def remove_paragraph_borders(style) -> None:
    paragraph_properties = style._element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15


def concern(doc: Document, index: int, heading: str, text: str, evidence: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(f"{index}. {heading}。").bold = True
    paragraph.add_run(f"{text} 证据位置：{evidence}。")


def markdown() -> str:
    lines = ["# DCS OPCT v11 BSPC 投稿前多视角审稿", "", "## 编辑信 Major revision", ""]
    lines.extend([EDITOR[0], "", EDITOR[1], ""])
    for reviewer in REVIEWERS:
        lines.extend([f"## {reviewer['role']}", "", "### 整体点评", "", reviewer["overall"], "", "### 主要问题", ""])
        for index, item in enumerate(reviewer["major"], 1):
            lines.extend([f"{index}. **{item[0]}**。{item[1]} 证据位置：{item[2]}。", ""])
        lines.extend(["### 次要问题", ""])
        for index, item in enumerate(reviewer["minor"], 1):
            lines.append(f"{index}. {item}")
        lines.extend(["", f"**置信度：{reviewer['confidence']}**", ""])
    lines.extend(["## 跨审稿意见汇总", ""])
    for cid, desc, raised, axis, rationale in CONSENSUS:
        lines.extend([f"**{cid} {desc}** | 提出者：{raised} | 轴：{axis} | 严重性：Major", "", rationale, ""])
    lines.extend(["## 修订任务表", "", "| ID | 审稿人 | 问题 | 策略 | 状态 | 所需输入 | 预期输出 | 是否阻塞 |", "|---|---|---|---|---|---|---|---|"])
    lines.extend("| " + " | ".join(row) + " |" for row in TASKS)
    lines.append("")
    return "\n".join(lines)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    for name, size in (("Normal", 10.5), ("Title", 16), ("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    remove_paragraph_borders(doc.styles["Title"])

    title = doc.add_paragraph("DCS OPCT v11 BSPC 投稿前多视角审稿", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("编辑信 Major revision", style="Heading 1")
    for item in EDITOR:
        body(doc, item)
    for reviewer in REVIEWERS:
        doc.add_paragraph(reviewer["role"], style="Heading 1")
        doc.add_paragraph("整体点评", style="Heading 2")
        body(doc, reviewer["overall"])
        doc.add_paragraph("主要问题", style="Heading 2")
        for index, item in enumerate(reviewer["major"], 1):
            concern(doc, index, *item)
        doc.add_paragraph("次要问题", style="Heading 2")
        for index, item in enumerate(reviewer["minor"], 1):
            body(doc, f"{index}. {item}")
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"置信度：{reviewer['confidence']}").bold = True

    doc.add_paragraph("跨审稿意见汇总", style="Heading 1")
    for cid, desc, raised, axis, rationale in CONSENSUS:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{cid} {desc}").bold = True
        paragraph.add_run(f" | 提出者：{raised} | 轴：{axis} | 严重性：Major")
        body(doc, rationale)

    doc.add_paragraph("修订任务表", style="Heading 1")
    headers = ("ID", "审稿人", "问题", "策略", "状态", "所需输入", "预期输出", "阻塞")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    repeat_header(table.rows[0])
    no_split(table.rows[0])
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    for row_index, values in enumerate(TASKS, 1):
        row = table.add_row()
        no_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                shade(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("nature-review-studio | 2026-09-14 | dcs-opct-v11-final")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.italic = True
    doc.save(DOCX)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    MARKDOWN.write_text(markdown(), encoding="utf-8")
    build_docx()
    print(DOCX)
    print(MARKDOWN)


if __name__ == "__main__":
    main()
